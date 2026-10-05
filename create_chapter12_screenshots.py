import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUTPUT_PATH = r"D:\libarr\LibraAI_Project_Screenshots.docx"
PDF_OUTPUT_PATH = r"D:\libarr\LibraAI_Project_Screenshots.pdf"
SCREENSHOT_DIR = r"D:\libarr\report_screenshots"

figures = [
    {
        "num": 1,
        "fig_label": "Figure 12.1",
        "file": "01_guest_home.png",
        "desc": "Figure 12.1: Guest Landing Page – Public discovery catalog with live database metrics, search bar, featured collections, and navigation."
    },
    {
        "num": 2,
        "fig_label": "Figure 12.2",
        "file": "02_login_popup.png",
        "desc": "Figure 12.2: User Authentication – Interactive modal dialog for student, faculty, librarian, and admin credential sign-in."
    },
    {
        "num": 3,
        "fig_label": "Figure 12.3",
        "file": "03_student_dashboard.png",
        "desc": "Figure 12.3: Student Dashboard – Active book loans with return due dates, instant renewal actions, and AI study cards."
    },
    {
        "num": 4,
        "fig_label": "Figure 12.4",
        "file": "04_catalog_search.png",
        "desc": "Figure 12.4: Books Catalog & Search – Real-time keyword filtering across 2,299 titles with rack location badges and availability counters."
    },
    {
        "num": 5,
        "fig_label": "Figure 12.5",
        "file": "05_ai_feature.png",
        "desc": "Figure 12.5: LibraBot AI Assistant – Domain-aware conversational assistant responding to technical queries with catalog references."
    },
    {
        "num": 6,
        "fig_label": "Figure 12.6",
        "file": "06_librarian_dashboard.png",
        "desc": "Figure 12.6: Librarian Dashboard – Daily circulation statistics, active loan monitors, pending returns, and rapid barcode actions."
    },
    {
        "num": 7,
        "fig_label": "Figure 12.7",
        "file": "07_issue_return.png",
        "desc": "Figure 12.7: Rapid Circulation Desk – Barcode circulation console with instant book issuing, return validation, and fine calculation."
    },
    {
        "num": 8,
        "fig_label": "Figure 12.8",
        "file": "08_admin_dashboard.png",
        "desc": "Figure 12.8: Admin Dashboard – Institutional KPI cards, 99.98% SLA monitoring, circulation activity curves, and AI telemetry."
    },
    {
        "num": 9,
        "fig_label": "Figure 12.9",
        "file": "09_admin_users.png",
        "desc": "Figure 12.9: Admin User Management – Campus member directory with role-based access control and account provisioning modal."
    },
    {
        "num": 10,
        "fig_label": "Figure 12.10",
        "file": "10_analytics_charts.png",
        "desc": "Figure 12.10: Analytics & Reports – Executive dashboard featuring circulation footfall timelines, category donuts, and top borrowed books."
    }
]

def set_cell_margins(cell, top=20, bottom=20, left=40, right=40):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, color="CBD5E1", sz="4", val="single"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_footer_page_number(run):
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def add_image_block(doc, fig_item, img_width_in=5.6):
    img_path = os.path.join(SCREENSHOT_DIR, fig_item["file"])
    if not os.path.exists(img_path):
        p_err = doc.add_paragraph(f"[Image Missing: {fig_item['file']}]")
        return

    # Bordered table frame for image
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Inches(img_width_in + 0.1)
    set_cell_margins(cell, top=20, bottom=20, left=20, right=20)
    set_cell_border(cell, color="CBD5E1", sz="4", val="single")

    p_img = cell.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(0)
    p_img.paragraph_format.space_after = Pt(0)
    r_img = p_img.add_run()
    r_img.add_picture(img_path, width=Inches(img_width_in))

    # Caption / Description directly below image
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2.5)
    p_cap.paragraph_format.space_after = Pt(8)
    p_cap.paragraph_format.line_spacing = 1.05
    r_cap = p_cap.add_run(fig_item["desc"])
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(30, 41, 59)

def generate_screenshots_document():
    doc = Document()

    # Page Margins: A4, 2 images per page layout
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)

        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_foot_label = p_foot.add_run("Page ")
        r_foot_label.font.name = "Times New Roman"
        r_foot_label.font.size = Pt(10)
        add_footer_page_number(p_foot.add_run())

    # =============================================================
    # PAGE 1: Header + Figure 1 & Figure 2
    # =============================================================
    # 1. "Chapter 12" in 16 pt font
    p_ch12 = doc.add_paragraph()
    p_ch12.paragraph_format.space_before = Pt(0)
    p_ch12.paragraph_format.space_after = Pt(1)
    p_ch12.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_ch12 = p_ch12.add_run("Chapter 12")
    r_ch12.font.name = "Times New Roman"
    r_ch12.font.size = Pt(16)
    r_ch12.font.bold = True
    r_ch12.font.color.rgb = RGBColor(0, 0, 0)

    # 2. "Output/Snapshots of Project" in 14 pt font
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(8)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_sub = p_sub.add_run("Output/Snapshots of Project")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0, 0, 0)

    # Image 1 (Figure 12.1)
    add_image_block(doc, figures[0], img_width_in=5.4)

    # Image 2 (Figure 12.2)
    add_image_block(doc, figures[1], img_width_in=5.4)

    # =============================================================
    # PAGE 2 to PAGE 5: 2 Images per page
    # =============================================================
    pairs = [
        (figures[2], figures[3]), # Page 2: Fig 3, Fig 4
        (figures[4], figures[5]), # Page 3: Fig 5, Fig 6
        (figures[6], figures[7]), # Page 4: Fig 7, Fig 8
        (figures[8], figures[9]), # Page 5: Fig 9, Fig 10
    ]

    for idx, (fig_a, fig_b) in enumerate(pairs):
        doc.add_page_break()
        add_image_block(doc, fig_a, img_width_in=5.5)
        add_image_block(doc, fig_b, img_width_in=5.5)

    doc.save(DOCX_OUTPUT_PATH)
    print(f"Successfully generated {DOCX_OUTPUT_PATH}")

if __name__ == "__main__":
    generate_screenshots_document()
